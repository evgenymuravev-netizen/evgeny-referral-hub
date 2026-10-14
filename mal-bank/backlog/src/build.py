import re
import sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from common import COLUMNS, TECH_COLS
from tab_cj import CJ
from tab_ops import UW, LP, RP
from tab_rf import RF
from tab_pl import PL

OUT = sys.argv[1]

# The non-platform tabs were written against an earlier platform numbering; map them once.
PL_MAP = {"PL-15": "PL-14", "PL-22": "PL-15", "PL-30": "PL-23", "PL-32": "PL-24"}


def remap(rows):
    rx = re.compile(r"\bPL-\d+\b")
    for r in rows:
        for i, v in enumerate(r):
            if isinstance(v, str):
                r[i] = rx.sub(lambda m: PL_MAP.get(m.group(0), m.group(0)), v)


for t in (CJ, UW, LP, RP, RF):
    remap(t)

TABS = [
    ("Customer journey", CJ, "Noor customer journey inside botim: entry, onboarding, Open Finance, pre-qualification, decision, offer, after payout, other products."),
    ("Underwriter", UW, "Mizan back office for underwriters, fraud analysts and collections: refer queue, case file, overrides, Decision log, monitoring cases."),
    ("Lending partner", LP, "What lending partners get from Noor: sharing agreements, the credit memo (UI, API, SFTP), callbacks, status feeds, partner reason codes."),
    ("Risk and policy", RP, "Risk policy owner, model risk and compliance: policy console, 4-eyes, simulation, reason-code library, models, dashboards."),
    ("Risk features", RF, "One row per risk feature Mizan must calculate from the data we have now, with the formula and parameters."),
    ("Platform and integrations", PL, "Integrations (botim, Open Finance, AECB, UAE PASS, vendors, partners) and the shared platform, security and operations work."),
]

# ---- checks: unique ids, every referenced id exists
ids = [r[0] for _, rows, _ in TABS for r in rows]
dupes = {i for i in ids if ids.count(i) > 1}
assert not dupes, dupes
known = set(ids)
missing = set()
for _, rows, _ in TABS:
    for r in rows:
        for v in r[1:]:
            if isinstance(v, str):
                for m in re.findall(r"\b(?:CJ|UW|LP|RP|RF|PL)-\d+\b", v):
                    if m not in known:
                        missing.add((r[0], m))
assert not missing, sorted(missing)

FONT = "Arial"
f_body = Font(name=FONT, size=10)
f_bold = Font(name=FONT, size=10, bold=True)
f_head = Font(name=FONT, size=10, bold=True, color="FFFFFF")
f_title = Font(name=FONT, size=14, bold=True)
fill_head = PatternFill("solid", fgColor="263238")
fill_tech_head = PatternFill("solid", fgColor="FFD966")
fill_tech = PatternFill("solid", fgColor="FFF2CC")
fill_band = PatternFill("solid", fgColor="ECEFF1")
thin = Side(style="thin", color="CFD8DC")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap_top = Alignment(wrap_text=True, vertical="top")
center_top = Alignment(horizontal="center", vertical="top", wrap_text=True)

REL = '"MVP,Phase 2,Later"'
SIZES = '"S,M,L,XL"'
STATUS = '"Not started,Estimating,Estimated,In progress,Done,Blocked,Dropped"'

wb = Workbook()
readme = wb.active
readme.title = "Read me"
ncols = len(COLUMNS)
first_tech = ncols - TECH_COLS + 1


def letter(i):
    from openpyxl.utils import get_column_letter
    return get_column_letter(i)


for name, rows, _ in TABS:
    ws = wb.create_sheet(name)
    for c, (h, w) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        tech = c >= first_tech
        cell.font = Font(name=FONT, size=10, bold=True, color="000000" if tech else "FFFFFF")
        cell.fill = fill_tech_head if tech else fill_head
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = border
        ws.column_dimensions[letter(c)].width = w
    for r_i, r in enumerate(rows, start=2):
        for c, v in enumerate(r, start=1):
            cell = ws.cell(row=r_i, column=c, value=v)
            cell.font = f_bold if c in (1, 2) else f_body
            cell.alignment = center_top if c in (7, 14, 17) else wrap_top
            cell.border = border
            if c >= first_tech:
                cell.fill = fill_tech
    last = len(rows) + 1
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{letter(ncols)}{last}"
    ws.row_dimensions[1].height = 30
    for formula, col in ((REL, 7), (SIZES, 14), (STATUS, ncols)):
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        dv.add(f"{letter(col)}2:{letter(col)}{last + 200}")
        ws.add_data_validation(dv)
    dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                        error="Enter person-days as a number, e.g. 8 or 2.5", errorTitle="Person-days")
    dv.add(f"{letter(first_tech)}2:{letter(first_tech)}{last + 200}")
    ws.add_data_validation(dv)

# ---- Read me
ws = readme


def para(row_i, text, bold_label=None):
    if bold_label is not None:
        ws.cell(row=row_i, column=1, value=bold_label).font = f_bold
        ws.cell(row=row_i, column=1).alignment = Alignment(vertical="top", wrap_text=True)
    ws.merge_cells(start_row=row_i, start_column=2, end_row=row_i, end_column=8)
    c = ws.cell(row=row_i, column=2, value=text)
    c.font = f_body
    c.alignment = Alignment(wrap_text=True, vertical="top")
    lines = max(1, -(-len(text) // 125))
    ws.row_dimensions[row_i].height = 14 * lines + 2

ws.column_dimensions["A"].width = 30
for col in "BCDEFG":
    ws.column_dimensions[col].width = 14
ws.column_dimensions["H"].width = 60
r = 1
ws.cell(row=r, column=1, value="Noor × Mizan product backlog").font = f_title
r += 1
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
c0 = ws.cell(row=r, column=1, value="Build backlog for the Noor customer journey inside botim and the Mizan back office. Written from the working prototype (v2.20) and its specs, so most rows point at a screen or rule that already exists there.")
c0.font = f_body
c0.alignment = Alignment(wrap_text=True, vertical="top")
ws.row_dimensions[r].height = 30
r += 2

ws.cell(row=r, column=1, value="How to estimate").font = f_bold
r += 1
howto = [
    "Fill the four yellow columns on each tab: Tech estimate (person-days, including tests and code review), Team / skills, Owner, Status.",
    "PM size is the product team's rough guess (S up to 3 person-days, M 4 to 10, L 11 to 20, XL more than 20). Keep it; it helps compare with your estimate.",
    "Built by says whose team builds it. Estimate only Noor's part; rows built by botim or the lender are listed so the dependencies are visible.",
    "Inputs and dependencies names the other rows a story needs (for example RF-17, PL-04). Risk features (RF) are separate tasks so they can be estimated and built on their own.",
    "Release assumes the MVP is the personal loan through botim with one lending partner (Partner Bank). Change the Release column if the scope changes.",
]
for t in howto:
    ws.cell(row=r, column=1, value="•").alignment = Alignment(horizontal="right", vertical="top")
    para(r, t)
    r += 1
r += 1

ws.cell(row=r, column=1, value="Example of the yellow columns (format only)").font = f_bold
r += 1
ex_head = ["ID", "Tech estimate (person-days)", "Team / skills", "Owner", "Status"]
ex_val = ["RF-17", 3, "BE 2, QA 1", "Backend lead", "Estimated"]
for c, h in enumerate(ex_head, start=1):
    cell = ws.cell(row=r, column=c, value=h)
    cell.font = Font(name=FONT, size=10, bold=True)
    cell.fill = fill_tech_head
    cell.border = border
    cell.alignment = Alignment(wrap_text=True, vertical="center")
ws.row_dimensions[r].height = 30
r += 1
for c, v in enumerate(ex_val, start=1):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font = Font(name=FONT, size=10, italic=True)
    cell.fill = fill_tech
    cell.border = border
r += 2

ws.cell(row=r, column=1, value="Summary by tab").font = f_bold
r += 1
heads = ["Tab", "Items", "MVP", "Phase 2", "Later", "Estimated items", "Person-days", "What it covers"]
for c, h in enumerate(heads, start=1):
    cell = ws.cell(row=r, column=c, value=h)
    cell.font = f_head
    cell.fill = fill_head
    cell.border = border
    cell.alignment = Alignment(wrap_text=True, vertical="center")
r += 1
top = r
for name, rows, desc in TABS:
    q = f"'{name}'"
    vals = [name, f"=COUNTA({q}!A:A)-1", f'=COUNTIF({q}!G:G,"MVP")', f'=COUNTIF({q}!G:G,"Phase 2")',
            f'=COUNTIF({q}!G:G,"Later")', f"=COUNT({q}!Q:Q)", f"=SUM({q}!Q:Q)", desc]
    for c, v in enumerate(vals, start=1):
        cell = ws.cell(row=r, column=c, value=v)
        cell.font = f_body
        cell.border = border
        cell.alignment = wrap_top if c in (1, 8) else Alignment(horizontal="center", vertical="top")
    r += 1
bottom = r - 1
ws.cell(row=r, column=1, value="Total").font = f_bold
for c in range(2, 8):
    cl = letter(c)
    cell = ws.cell(row=r, column=c, value=f"=SUM({cl}{top}:{cl}{bottom})")
    cell.font = f_bold
    cell.alignment = Alignment(horizontal="center")
for c in range(1, 9):
    ws.cell(row=r, column=c).border = border
    ws.cell(row=r, column=c).fill = fill_band
r += 2

ws.cell(row=r, column=1, value="Columns").font = f_bold
r += 1
legend = [
    ("ID", "CJ customer journey, UW underwriter, LP lending partner, RP risk and policy, RF risk feature, PL platform and integrations."),
    ("Type", "User story, Feature, Risk feature, Integration, Tech enabler or Compliance."),
    ("Actor", "Who the story is for (customer, underwriter, Partner Bank credit officer, risk policy owner, system)."),
    ("Built by", "Noor, botim, Partner Bank or a combination. Estimate Noor's share."),
    ("Release", "MVP: first funded personal loans through botim with Partner Bank. Phase 2: monitoring, car loan, split, newcomer evidence, second lender, data pipelines for the model. Later: SME, starter and upgrade, salary advance, trained model, NoorScore v2."),
    ("Business rules / calculation", "For risk features: the formula, parameters and their bounds as in the prototype engine. Parameter values are illustrative until the risk team signs them off."),
    ("Prototype reference", "Screen and step in the prototype, the engine function, or the spec file in the repository."),
    ("Components", "FE front end, BE back end (Mizan services), Data (pipelines, features), ML, Integration, Ops, QA."),
]
for k, v in legend:
    para(r, v, k)
    r += 1
r += 1

ws.cell(row=r, column=1, value="Rules that apply to every item").font = f_bold
r += 1
rules = [
    "Nationality, sex, religion, health, family status and the destination of a transfer or trip are never score, limit or price inputs. Nationality from the Emirates ID only routes the newcomer evidence offer.",
    "Red flags and data gaps can only refer a case to a person at origination: never an automatic decline, never a price change.",
    "Lenders receive scores, reason codes and banded aggregates only. Methodology, thresholds, band boundaries, data-source mix, raw data and Emirates ID images stay inside Noor.",
    "Reason codes have three colours only: green, yellow, red. Customers see the RC_* adverse-action reasons in English and Arabic; lenders see library codes.",
    "Regulatory caps (DBR 50%, retirees 30%, 20× salary, 48 months, car LTV 80% and 60 months) are enforced in the engine whatever the policy says.",
    "Every policy, library or agreement change needs a second approver (4-eyes). Every decision records the policy, model and engine versions and can be reproduced.",
]
for t in rules:
    ws.cell(row=r, column=1, value="•").alignment = Alignment(horizontal="right", vertical="top")
    para(r, t)
    r += 1
r += 1

ws.cell(row=r, column=1, value="Sources").font = f_bold
r += 1
src = [
    ("Prototype", "https://evgenymuravev-netizen.github.io/evgeny-referral-hub/noor.html"),
    ("Lender view", "https://evgenymuravev-netizen.github.io/evgeny-referral-hub/noor.html#lender"),
    ("Specs and engine", "https://github.com/evgenymuravev-netizen/evgeny-referral-hub/tree/main/mal-bank/prototype"),
]
for k, v in src:
    ws.cell(row=r, column=1, value=k).font = f_bold
    c = ws.cell(row=r, column=2, value=v)
    c.hyperlink = v
    c.font = Font(name=FONT, size=10, color="1155CC", underline="single")
    r += 1
ws.freeze_panes = None

# ---- Glossary
gl = wb.create_sheet("Glossary")
gl.column_dimensions["A"].width = 26
gl.column_dimensions["B"].width = 110
for c, h in enumerate(["Term", "Meaning"], start=1):
    cell = gl.cell(row=1, column=c, value=h)
    cell.font = f_head
    cell.fill = fill_head
terms = [
    ("Noor", "The customer-facing credit product, embedded in botim."),
    ("Mizan", "Noor's decision engine and back office (Workbench, Decision log, Policy console, Lender view, Monitoring)."),
    ("NoorScore", "Noor's credit score. Phase 1 is an expert scorecard on top of the bureau score; Phase 2 is a model trained on repayment data."),
    ("Data completeness", "Step 1 of the score: how much of the customer's financial picture Noor sees (current account, salary, cards, loans), 0 to 100%."),
    ("AECB", "Al Etihad Credit Bureau, the UAE credit bureau. Reports are pulled through the licensed lender acting for Noor."),
    ("Open Finance, Al Tareq, Nebras", "UAE Open Finance framework (CBUAE). Al Tareq is the consumer-facing brand; Nebras runs the platform. AIS reads account data; PIS initiates payments."),
    ("UAE PASS", "The national digital identity, used for strong identity and digital signatures."),
    ("Emirates ID (EID)", "The national ID card. Noor scans the front and back."),
    ("WPS", "Wage Protection System: salaries paid through it carry a WPS marker."),
    ("DBR", "Debt burden ratio: monthly debt payments (including the new instalment) divided by salary plus regular income. CBUAE cap 50% (Regulation 29/2011)."),
    ("ESR", "Expense-to-salary ratio on the bureau file: bureau instalments divided by salary."),
    ("FCF", "Free cash flow: income minus spending minus debt payments, per month."),
    ("LTV", "Loan-to-value for the car loan: loan divided by vehicle price, capped at 80%."),
    ("DSCR, NOCF", "SME: net operating cash flow divided by the instalment; NOCF is average revenue minus average operating outflows."),
    ("KFS", "Key Facts Statement, in Arabic and English, acknowledged before signing."),
    ("Cooling-off", "5 business days after signing before payout, or a written waiver (CBUAE Consumer Protection Regulation 8/2020)."),
    ("PDPL", "UAE Personal Data Protection Law."),
    ("DPD", "Days past due on a loan (DPD30, DPD90)."),
    ("Refer", "Outcome that sends a case to a person (underwriter or fraud analyst) instead of deciding automatically."),
    ("Red flag", "A change against the customer's own history (cards maxed, remittance surge, balances drained, travel after payout, payments stopped)."),
    ("Reason-code library", "The codes lenders see (G01 to G11, Y01 to Y14, R01 to R09, plus partner codes). Customers see RC_* reasons."),
    ("Sharing agreement", "Versioned list of what each lender receives in the credit memo."),
    ("Credit memo", "Noor's summary of a decision for the lender: NoorScore and band, reason codes, completeness and banded aggregates."),
    ("4-eyes", "A change needs a second person to approve it; the approver cannot be the author."),
    ("SHADOW and LIVE", "SHADOW: computed and logged but not applied. LIVE: applied to decisions."),
    ("SHAP", "Method to show how much each feature moved a model's output; used for Phase 2 explanations."),
    ("Credit Passport", "Import of a home-country bureau file with the customer's consent (newcomers)."),
    ("pickMin", "Approved amount is the lowest of all caps (request, DBR, free cash flow, salary multiple, product cap and product-specific caps)."),
]
for i, (k, v) in enumerate(terms, start=2):
    gl.cell(row=i, column=1, value=k).font = f_bold
    c = gl.cell(row=i, column=2, value=v)
    c.font = f_body
    c.alignment = Alignment(wrap_text=True, vertical="top")
gl.freeze_panes = "A2"

wb.save(OUT)
print("rows:", {n: len(rw) for n, rw, _ in TABS}, "total", sum(len(rw) for _, rw, _ in TABS))
